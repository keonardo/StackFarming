"""Generate 卡牌数据总表.xlsx from StackFarming game data."""
import os, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ── Styles ──────────────────────────────────────────────
header_font = Font(name="SimHei", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
title_font = Font(name="SimHei", size=14, bold=True, color="1F4E79")
data_font = Font(name="SimHei", size=10)
note_font = Font(name="SimHei", size=9, italic=True, color="666666")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"))
center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
alt_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")


def style_header(ws, row, cols):
    for col in range(1, cols + 1):
        c = ws.cell(row=row, column=col)
        c.font = header_font; c.fill = header_fill
        c.alignment = center_align; c.border = thin_border


def style_row(ws, row, cols, align=center_align):
    for col in range(1, cols + 1):
        c = ws.cell(row=row, column=col)
        c.font = data_font; c.border = thin_border; c.alignment = align


def auto_width(ws, min_w=8, max_w=40):
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=0)
        width = min(max(max_len * 1.3 + 2, min_w), max_w)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width


def add_title(ws, text, row=1, merge="A1:L1"):
    ws.merge_cells(merge)
    ws["A1"] = text
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 30


# ═══════════════════════════════════════════════════════════
# Sheet 1: 卡牌数据总表
# ═══════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "卡牌数据总表"

add_title(ws1, "StackFarming 卡牌数值总表", merge="A1:L1")
ws1.merge_cells("A2:L2")
ws1["A2"] = ("生物(Creature): ValueA=Health, ValueB=Progress  |  "
             "资源(Resource): ValueA=Intensity, ValueB=Duration  |  "
             "地貌(Container): ValueA=Capacity, ValueB=Moisture")
ws1["A2"].font = note_font
ws1["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws1.row_dimensions[2].height = 25

h1 = ["卡牌名称", "英文名", "类别(CardRole)", "类型(CardType)", "Type值",
      "ValueA初始", "ValueA含义", "ValueB初始", "ValueB含义",
      "特殊属性", "交互逻辑", "死亡/消失"]
for i, h in enumerate(h1, 1):
    ws1.cell(row=3, column=i, value=h)
style_header(ws1, 3, len(h1))
ws1.row_dimensions[3].height = 35

cards = [
    # ── 生物卡 ──
    ["小鸭子", "Duckling", "生物", "动物(Animal)", 0,
     100, "健康值(Health)", 50, "生长进度(Progress)",
     "代谢率10/s", "捕食害虫→重置Progress+产粪; 超载时背叛扣作物HP", "→枯草"],
    ["水稻", "Rice", "生物", "作物(Crop)", 1,
     100, "健康值(Health)", 0, "生长进度(Progress)",
     "基础生长5%/s, 粪便倍率1.5x", "粪便堆叠加速生长→满Progress产稻谷", "→枯草"],
    ["鱼苗", "Fish Fry", "生物", "鱼类(Fish)", 2,
     100, "健康值(Health)", 0, "生长进度(Progress)",
     "水生生物", "滋润度=0时失血15HP/s; 翻塘时瞬间死亡", "→死鱼"],
    ["害虫", "Pest", "生物", "害虫(Pest)", 6,
     50, "健康值(Health)", 0, "—",
     "攻击5HP/s, 间隔1s", "不参与堆叠; 附着作物扣血; 可被动物捕食", "→无(直接销毁)"],

    # ── 资源卡 ──
    ["粪便卡", "Feces", "资源", "粪便(Feces)", 3,
     10, "强度(Intensity)", 60, "持续时间(Duration)s",
     "Duration倒计时60s后自毁", "堆叠作物→加速生长; 过多+无作物→翻塘; 可漂移到下游", "→无(过期消失)"],
    ["死鱼", "Dead Fish", "资源", "死鱼(DeadFish)", 4,
     10, "强度(Intensity)", 60, "持续时间(Duration)s",
     "默认60s存续", "鱼类死亡转化; 可被水流漂移", "→无(过期消失)"],
    ["枯草", "Dry Grass", "资源", "枯草(DryGrass)", 5,
     10, "强度(Intensity)", 60, "持续时间(Duration)s",
     "默认60s存续", "动植物死亡转化; 稻谷收获产出", "→无(过期消失)"],

    # ── 地貌卡 ──
    ["水稻田", "Rice Field", "地貌", "农田(Farmland)", 7,
     4, "承载力(Capacity)", 100, "滋润度(Moisture)",
     "干燥失血15HP/s", "承载作物/动物; 通过水渠灌溉; 惊蛰时生成3-5害虫", "—"],
    ["旱田", "Dry Field", "地貌", "农田(Farmland)", 7,
     4, "承载力(Capacity)", 100, "滋润度(Moisture)",
     "同水稻田", "同水稻田(无水稻标记, 不受惊蛰害虫影响)", "—"],
    ["深水鱼凼", "Deep Water Pond", "地貌", "深水鱼凼(DeepWaterFishPond)", 8,
     5, "承载力(Capacity)", 100, "滋润度(Moisture)",
     "芒种大旱时保留灌溉", "芒种时不受干旱影响; 专门养鱼", "—"],
    ["河道", "River", "地貌", "河道(River)", 9,
     999, "承载力(极大)", 100, "滋润度(Moisture)",
     "初始已灌溉", "水源节点; BFS灌溉传播起点", "—"],
    ["水渠", "Canal", "地貌", "水渠(Canal)", 10,
     1, "承载力(极小)", 100, "滋润度(Moisture)",
     "连接距离160px", "中继节点; 连接河道与农田传递灌溉", "—"],
]

for i, row_data in enumerate(cards):
    r = i + 4
    for j, val in enumerate(row_data):
        ws1.cell(row=r, column=j + 1, value=val)
    style_row(ws1, r, len(h1))
    if i % 2 == 0:
        for col in range(1, len(h1) + 1):
            ws1.cell(row=r, column=col).fill = alt_fill

auto_width(ws1, min_w=10, max_w=38)
ws1.column_dimensions['A'].width = 12
ws1.column_dimensions['K'].width = 40
ws1.column_dimensions['L'].width = 22

# ═══════════════════════════════════════════════════════════
# Sheet 2: 核心计算公式
# ═══════════════════════════════════════════════════════════
ws2 = wb.create_sheet("核心计算公式")
add_title(ws2, "核心交互公式 (from CLAUDE.md)", merge="A1:D1")

formulas = [
    ["ID", "名称", "计算逻辑", "代码位置"],
    ["F1", "干燥失血",
     "!IsIrrigated => 水生Bio.Health -= 15 * Delta",
     "container_card.gd: _check_survival_and_imbalance()"],
    ["F2", "作物生长",
     "Crop.Progress += (5 + Feces.Intensity * 1.5) * SolarTermModifier * Delta",
     "creature_card.gd: _process_crop_growth()"],
    ["F3", "捕食",
     "Animal overlaps Pest => Pest.queue_free(), Animal.Progress=0, SpawnFeces()",
     "creature_card.gd: _process_animal_interactions()"],
    ["F4", "代谢产粪",
     "Animal.Progress >= 100 => SpawnFeces(), Progress=0",
     "creature_card.gd: _process_animal_metabolism()"],
    ["F5", "翻塘",
     "FecesCount >= 3 && CropCount == 0 => All Fish.Health = 0",
     "container_card.gd: _check_survival_and_imbalance()"],
    ["F6", "背叛扣血",
     "AnimalCount > Capacity => Crop.Health -= 8 * Delta",
     "creature_card.gd: _process_animal_interactions()"],
    ["F7", "灌溉BFS",
     "BFS(Rivers→Canals→Containers); distance<=160px => IsIrrigated=true",
     "irrigation_manager.gd: check_irrigation_linkage()"],
    ["F8", "粪便漂移",
     "FecesCount>1且下游Y>上游Y+25 => 顶部粪便漂移到下游地块",
     "irrigation_manager.gd: _process_substance_drift()"],
    ["F9", "霜降冻伤",
     "暴露热带生物.Health -= 10 * Delta",
     "solar_term_engine.gd: _process_shuangjiang_frost_damage()"],
    ["F10", "死亡转化",
     "Health<=0 => Animal/Crop→枯草, Fish→死鱼; 继承堆叠关系",
     "base_card.gd: on_health_zero()"],
]

for i, row_data in enumerate(formulas):
    for j, val in enumerate(row_data):
        ws2.cell(row=i + 2, column=j + 1, value=val)
style_header(ws2, 2, len(formulas[0]))
for i in range(3, len(formulas) + 2):
    style_row(ws2, i, len(formulas[0]), align=left_align)
auto_width(ws2, min_w=12, max_w=60)
ws2.column_dimensions['C'].width = 55
ws2.column_dimensions['D'].width = 45

# ═══════════════════════════════════════════════════════════
# Sheet 3: 二十四节气引擎
# ═══════════════════════════════════════════════════════════
ws3 = wb.create_sheet("二十四节气引擎")
add_title(ws3, "Solar Term Engine — 二十四节气", merge="A1:E1")

terms = [
    ["节气", "枚举值", "轮转周期", "全局Buff", "突发事件"],
    ["无(None)", 0, "term_duration\n(编辑器120s/场景10s)", "无", "—"],
    ["惊蛰(Jingzhe)", 1, "同上", "—",
     "所有水稻田刷3-5张害虫卡(随机分布±20px)"],
    ["芒种(Mangzhong)", 2, "同上", "—",
     "大旱: 非深水鱼凼地块IsIrrigated=false; 触发灌溉重检"],
    ["霜降(Shuangjiang)", 3, "同上", "CropGrowth*2.0",
     "暴露热带生物失血10HP/s; 节气结束后生长率恢复1.0"],
]

for i, row_data in enumerate(terms):
    for j, val in enumerate(row_data):
        ws3.cell(row=i + 2, column=j + 1, value=val)
style_header(ws3, 2, len(terms[0]))
for i in range(3, len(terms) + 2):
    style_row(ws3, i, len(terms[0]), align=left_align)
auto_width(ws3, min_w=12, max_w=55)
ws3.column_dimensions['E'].width = 50

# ═══════════════════════════════════════════════════════════
# Sheet 4: 默认参数速查
# ═══════════════════════════════════════════════════════════
ws4 = wb.create_sheet("默认参数速查")
add_title(ws4, "全局默认参数 (可在Godot编辑器@export调整)", merge="A1:C1")

params = [
    ["参数名", "默认值", "所属脚本"],
    ["MetabolismRate", "10/s", "creature_card.gd"],
    ["BaseCropGrowthRate", "5%/s", "creature_card.gd"],
    ["FecesMultiplier", "1.5x", "creature_card.gd"],
    ["OverCapacityCropDamage", "8HP/s", "creature_card.gd"],
    ["PreyCheckInterval", "1.0s", "creature_card.gd"],
    ["DryHealthDecayPerSecond", "15HP/s", "container_card.gd"],
    ["FecesPondTurningThreshold", "3", "container_card.gd"],
    ["DamagePerSecond (害虫)", "5HP/s", "pest_card.gd"],
    ["TargetSearchInterval (害虫)", "1.0s", "pest_card.gd"],
    ["ConnectionDistance (灌溉)", "160px", "irrigation_manager.gd"],
    ["DriftDistance (漂移)", "250px", "irrigation_manager.gd"],
    ["DriftInterval", "5.0s", "irrigation_manager.gd"],
    ["LinkageCheckInterval", "1.0s", "irrigation_manager.gd"],
    ["TermDuration", "120s(编辑器)/10s(场景覆写)", "solar_term_engine.gd"],
    ["资源卡默认Intensity", "10", "card_spawner.gd"],
    ["资源卡默认Duration", "60s", "card_spawner.gd"],
    ["生物卡默认Health", "100", "card_spawner.gd"],
    ["生物卡默认Progress", "0", "card_spawner.gd"],
    ["容器卡默认Capacity", "5", "card_spawner.gd"],
    ["容器卡默认Moisture", "100", "card_spawner.gd"],
    ["害虫卡默认Health", "50", "card_spawner.gd"],
]

for i, row_data in enumerate(params):
    for j, val in enumerate(row_data):
        ws4.cell(row=i + 2, column=j + 1, value=val)
style_header(ws4, 2, len(params[0]))
for i in range(3, len(params) + 2):
    style_row(ws4, i, len(params[0]), align=left_align)
auto_width(ws4, min_w=16, max_w=45)

# ═══════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════
output_dir = os.path.dirname(os.path.abspath(__file__))
output_path = os.path.join(os.path.dirname(output_dir), "docs", "卡牌数据总表.xlsx")
os.makedirs(os.path.dirname(output_path), exist_ok=True)
wb.save(output_path)
print(f"OK — saved to: {output_path}")
print(f"Sheets: {wb.sheetnames}")
